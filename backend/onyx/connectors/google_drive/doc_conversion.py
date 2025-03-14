import io
from datetime import datetime
from datetime import timezone
from tempfile import NamedTemporaryFile

import openpyxl  # type: ignore
from googleapiclient.discovery import build  # type: ignore
from googleapiclient.errors import HttpError  # type: ignore

from onyx.configs.app_configs import CONTINUE_ON_CONNECTOR_FAILURE
from onyx.configs.constants import DocumentSource
from onyx.configs.constants import IGNORE_FOR_QA
from onyx.connectors.google_drive.constants import DRIVE_FOLDER_TYPE
from onyx.connectors.google_drive.constants import DRIVE_SHORTCUT_TYPE
from onyx.connectors.google_drive.constants import UNSUPPORTED_FILE_TYPE_CONTENT
from onyx.connectors.google_drive.models import GDriveMimeType
from onyx.connectors.google_drive.models import GoogleDriveFileType
from onyx.connectors.google_drive.section_extraction import get_document_sections
from onyx.connectors.google_utils.resources import GoogleDocsService
from onyx.connectors.google_utils.resources import GoogleDriveService
from onyx.connectors.models import Document
from onyx.connectors.models import Section
from onyx.connectors.models import SlimDocument
from onyx.file_processing.extract_file_text import docx_to_text
from onyx.file_processing.extract_file_text import pptx_to_text
from onyx.file_processing.extract_file_text import read_pdf_file
from onyx.file_processing.unstructured import get_unstructured_api_key
from onyx.file_processing.unstructured import unstructured_to_text
from onyx.utils.logger import setup_logger

logger = setup_logger()


# these errors don't represent a failure in the connector, but simply files
# that can't / shouldn't be indexed
ERRORS_TO_CONTINUE_ON = [
    "cannotExportFile",
    "exportSizeLimitExceeded",
    "cannotDownloadFile",
]


def _extract_sections_basic(
    file: dict[str, str], service: GoogleDriveService
) -> list[Section]:
    mime_type = file["mimeType"]
    link = file["webViewLink"]
    supported_file_types = set(item.value for item in GDriveMimeType)

    if mime_type not in supported_file_types:
        # Unsupported file types can still have a title, finding this way is still useful
        return [Section(link=link, text=UNSUPPORTED_FILE_TYPE_CONTENT)]

    try:
        # ---------------------------
        # Google Sheets extraction
        if mime_type == GDriveMimeType.SPREADSHEET.value:
            try:
                sheets_service = build(
                    "sheets", "v4", credentials=service._http.credentials
                )
                spreadsheet = (
                    sheets_service.spreadsheets()
                    .get(spreadsheetId=file["id"])
                    .execute()
                )

                sections = []
                for sheet in spreadsheet["sheets"]:
                    sheet_name = sheet["properties"]["title"]
                    sheet_id = sheet["properties"]["sheetId"]

                    # Get sheet dimensions
                    grid_properties = sheet["properties"].get("gridProperties", {})
                    row_count = grid_properties.get("rowCount", 1000)
                    column_count = grid_properties.get("columnCount", 26)

                    # Convert column count to letter (e.g., 26 -> Z, 27 -> AA)
                    end_column = ""
                    while column_count:
                        column_count, remainder = divmod(column_count - 1, 26)
                        end_column = chr(65 + remainder) + end_column

                    range_name = f"'{sheet_name}'!A1:{end_column}{row_count}"

                    try:
                        result = (
                            sheets_service.spreadsheets()
                            .values()
                            .get(spreadsheetId=file["id"], range=range_name)
                            .execute()
                        )
                        values = result.get("values", [])

                        if values:
                            text = f"Sheet: {sheet_name}\n"
                            for row in values:
                                text += "\t".join(str(cell) for cell in row) + "\n"
                            sections.append(
                                Section(
                                    link=f"{link}#gid={sheet_id}",
                                    text=text,
                                )
                            )
                    except HttpError as e:
                        logger.warning(
                            f"Error fetching data for sheet '{sheet_name}': {e}"
                        )
                        continue
                return sections

            except Exception as e:
                logger.warning(
                    f"Ran into exception '{e}' when pulling data from Google Sheet '{file['name']}'."
                    " Falling back to basic extraction."
                )
        # ---------------------------
        # Microsoft Excel (.xlsx or .xls) extraction branch
        elif mime_type in [
            GDriveMimeType.SPREADSHEET_OPEN_FORMAT.value,
            GDriveMimeType.SPREADSHEET_MS_EXCEL.value,
        ]:
            try:
                response = service.files().get_media(fileId=file["id"]).execute()

                with NamedTemporaryFile(suffix=".xlsx", delete=True) as tmp:
                    tmp.write(response)
                    tmp_path = tmp.name

                    section_separator = "\n\n"
                    workbook = openpyxl.load_workbook(tmp_path, read_only=True)

                    # Work similarly to the xlsx_to_text function used for file connector
                    # but returns Sections instead of a string
                    sections = [
                        Section(
                            link=link,
                            text=(
                                f"Sheet: {sheet.title}\n\n"
                                + section_separator.join(
                                    ",".join(map(str, row))
                                    for row in sheet.iter_rows(
                                        min_row=1, values_only=True
                                    )
                                    if row
                                )
                            ),
                        )
                        for sheet in workbook.worksheets
                    ]

                return sections

            except Exception as e:
                logger.warning(
                    f"Error extracting data from Excel file '{file['name']}': {e}"
                )
                return [
                    Section(link=link, text="Error extracting data from Excel file")
                ]

        # ---------------------------
        # Export for Google Docs, PPT, and fallback for spreadsheets
        if mime_type in [
            GDriveMimeType.DOC.value,
            GDriveMimeType.PPT.value,
            GDriveMimeType.SPREADSHEET.value,
        ]:
            export_mime_type = (
                "text/plain"
                if mime_type != GDriveMimeType.SPREADSHEET.value
                else "text/csv"
            )
            text = (
                service.files()
                .export(fileId=file["id"], mimeType=export_mime_type)
                .execute()
                .decode("utf-8")
            )
            return [Section(link=link, text=text)]

        # ---------------------------
        # Plain text and Markdown files
        elif mime_type in [
            GDriveMimeType.PLAIN_TEXT.value,
            GDriveMimeType.MARKDOWN.value,
        ]:
            return [
                Section(
                    link=link,
                    text=service.files()
                    .get_media(fileId=file["id"])
                    .execute()
                    .decode("utf-8"),
                )
            ]
        # ---------------------------
        # Word, PowerPoint, PDF files
        if mime_type in [
            GDriveMimeType.WORD_DOC.value,
            GDriveMimeType.POWERPOINT.value,
            GDriveMimeType.PDF.value,
        ]:
            response = service.files().get_media(fileId=file["id"]).execute()
            if get_unstructured_api_key():
                return [
                    Section(
                        link=link,
                        text=unstructured_to_text(
                            file=io.BytesIO(response),
                            file_name=file.get("name", file["id"]),
                        ),
                    )
                ]

            if mime_type == GDriveMimeType.WORD_DOC.value:
                return [
                    Section(link=link, text=docx_to_text(file=io.BytesIO(response)))
                ]
            elif mime_type == GDriveMimeType.PDF.value:
                text, _ = read_pdf_file(file=io.BytesIO(response))
                return [Section(link=link, text=text)]
            elif mime_type == GDriveMimeType.POWERPOINT.value:
                return [
                    Section(link=link, text=pptx_to_text(file=io.BytesIO(response)))
                ]

        # Catch-all case, should not happen since there should be specific handling
        # for each of the supported file types
        error_message = f"Unsupported file type: {mime_type}"
        logger.error(error_message)
        raise ValueError(error_message)

    except Exception:
        return [Section(link=link, text=UNSUPPORTED_FILE_TYPE_CONTENT)]


def convert_drive_item_to_document(
    file: GoogleDriveFileType,
    drive_service: GoogleDriveService,
    docs_service: GoogleDocsService,
) -> Document | None:
    try:
        # Skip files that are shortcuts
        if file.get("mimeType") == DRIVE_SHORTCUT_TYPE:
            logger.info("Ignoring Drive Shortcut Filetype")
            return None
        # Skip files that are folders
        if file.get("mimeType") == DRIVE_FOLDER_TYPE:
            logger.info("Ignoring Drive Folder Filetype")
            return None

        sections: list[Section] = []

        # Special handling for Google Docs to preserve structure, link
        # to headers
        if file.get("mimeType") == GDriveMimeType.DOC.value:
            try:
                sections = get_document_sections(docs_service, file["id"])
            except Exception as e:
                logger.warning(
                    f"Ran into exception '{e}' when pulling sections from Google Doc '{file['name']}'."
                    " Falling back to basic extraction."
                )
        # NOTE: this will run for either (1) the above failed or (2) the file is not a Google Doc
        if not sections:
            try:
                # For all other file types just extract the text
                sections = _extract_sections_basic(file, drive_service)

            except HttpError as e:
                reason = e.error_details[0]["reason"] if e.error_details else e.reason
                message = e.error_details[0]["message"] if e.error_details else e.reason
                if e.status_code == 403 and reason in ERRORS_TO_CONTINUE_ON:
                    logger.warning(
                        f"Could not export file '{file['name']}' due to '{message}', skipping..."
                    )
                    return None

                raise
        if not sections:
            return None

        return Document(
            id=file["webViewLink"],
            sections=sections,
            source=DocumentSource.GOOGLE_DRIVE,
            semantic_identifier=file["name"],
            doc_updated_at=datetime.fromisoformat(file["modifiedTime"]).astimezone(
                timezone.utc
            ),
            metadata={}
            if any(section.text for section in sections)
            else {IGNORE_FOR_QA: "True"},
            additional_info=file.get("id"),
        )
    except Exception as e:
        if not CONTINUE_ON_CONNECTOR_FAILURE:
            raise e

        logger.exception("Ran into exception when pulling a file from Google Drive")
    return None


def build_slim_document(file: GoogleDriveFileType) -> SlimDocument | None:
    # Skip files that are folders or shortcuts
    if file.get("mimeType") in [DRIVE_FOLDER_TYPE, DRIVE_SHORTCUT_TYPE]:
        return None

    return SlimDocument(
        id=file["webViewLink"],
        perm_sync_data={
            "doc_id": file.get("id"),
            "drive_id": file.get("driveId"),
            "permissions": file.get("permissions", []),
            "permission_ids": file.get("permissionIds", []),
            "name": file.get("name"),
            "owner_email": file.get("owners", [{}])[0].get("emailAddress"),
        },
    )
